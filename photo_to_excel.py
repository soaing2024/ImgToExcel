import time

from PIL import Image
from openpyxl.styles import PatternFill
import openpyxl
from openpyxl.styles import PatternFill
from openpyxl import load_workbook


def number_to_excel_column(n):
    """将数字转换为Excel风格的列字母组合

    参数:
        n (int): 要转换的数字（从1开始）

    返回:
        str: 对应的字母组合（A-Z, AA-AZ, BA-BZ, ..., AAA等）
    """
    if n < 1:
        raise ValueError("输入数字必须大于等于1")

    result = []
    while n > 0:
        n -= 1  # 调整为0-25的范围
        remainder = n % 26
        result.append(chr(65 + remainder))  # 65是'A'的ASCII码
        n = n // 26

    return ''.join(reversed(result))


def rgb_to_hex(r, g, b):
    # 处理每个分量：四舍五入并限制在0-255范围内
    r = max(0, min(255, round(r)))
    g = max(0, min(255, round(g)))
    b = max(0, min(255, round(b)))

    # 格式化为两位十六进制，大写字母，拼接成字符串
    return '{:02X}{:02X}{:02X}'.format(r, g, b)


# 比如打开test.xlsx
wb = load_workbook(filename='1.xlsx')
# 使用第一个sheet作为工作簿
work = wb[wb.sheetnames[0]]


image = Image.open(r"C:\Users\开摆的林十三\Downloads\生成二次元头像.png")

width, height = image.size

# 调整行高
# work.row_dimensions['1'].height = 10
num = 0

for y in range(int(height/3)):
    work.row_dimensions[y+1].height = 5
    for x in range(int(width/3)):
        # 获取像素的RGB值
        color = image.getpixel((x*3, y*3))
        print(x, y)

        if num == 0:
            pass
            work.column_dimensions[f'{number_to_excel_column(x+1)}'].width = 1

        r, g, b = color[0], color[1], color[2]

        # 打印像素的RGB值
        print(f"像素位置: ({number_to_excel_column(x+1)}, {y+1}), RGB值: {rgb_to_hex(r=r, g=g, b=b)}")

        d4 = work[f'{number_to_excel_column(x+1)}{y+1}']
        d4.fill = PatternFill('solid', fgColor=rgb_to_hex(r=r, g=g, b=b))

    num += 1


wb.close()
wb.save('result.xlsx')

